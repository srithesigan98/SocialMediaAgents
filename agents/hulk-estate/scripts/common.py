#!/usr/bin/env python3
"""Shared helpers for Hulk Estate: config, listings loading, WhatsApp links, run state."""
from __future__ import annotations

import csv
import io
import json
import os
import re
import urllib.parse
from datetime import datetime, timedelta, timezone
from pathlib import Path

import requests
import yaml

HERE = Path(__file__).resolve().parent
AGENT_DIR = HERE.parent
CONFIG_PATH = AGENT_DIR / "config" / "agent.yaml"
STATE_DIR = HERE / "state"


def load_config() -> dict:
    cfg = yaml.safe_load(CONFIG_PATH.read_text(encoding="utf-8"))
    # Secrets / private URLs may override the file so they never land in git.
    if os.environ.get("LISTINGS_URL"):
        cfg["listings"]["url"] = os.environ["LISTINGS_URL"]
    if os.environ.get("WHATSAPP_NUMBER"):
        cfg["whatsapp"]["number"] = os.environ["WHATSAPP_NUMBER"]
    return cfg


# --------------------------------------------------------------------------- listings

def _norm(header: str) -> str:
    return re.sub(r"[^a-z0-9]", "", header.lower())


def _rows_from_csv_text(text: str) -> list[dict]:
    return list(csv.DictReader(io.StringIO(text)))


def _rows_from_xlsx(path: Path) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover - depends on optional extra
        raise SystemExit("source: xlsx needs openpyxl — run: pip install openpyxl")
    ws = load_workbook(path, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h or "") for h in rows[0]]
    return [
        {h: ("" if v is None else str(v)) for h, v in zip(headers, r)}
        for r in rows[1:]
        if any(v is not None and str(v).strip() for v in r)
    ]


def _resolve_path(raw: str) -> Path:
    """Listing paths in agent.yaml are relative to the config/ directory."""
    path = Path(raw)
    return path if path.is_absolute() else (CONFIG_PATH.parent / path).resolve()


def _gsheet_csv_url(url: str) -> str:
    """Accept a normal Google Sheets share URL and turn it into a CSV export URL."""
    m = re.search(r"/spreadsheets/d/([A-Za-z0-9_-]+)", url)
    if not m:
        return url  # already a published-CSV or other direct link
    doc_id = m.group(1)
    gid_match = re.search(r"[#&?]gid=(\d+)", url)
    gid = gid_match.group(1) if gid_match else "0"
    return f"https://docs.google.com/spreadsheets/d/{doc_id}/export?format=csv&gid={gid}"


def load_listings(cfg: dict) -> list[dict]:
    """Return listings as dicts keyed by the AGENT's field names (see config columns map)."""
    lc = cfg["listings"]
    source = lc["source"]

    if source == "gsheet":
        url = lc.get("url") or ""
        if not url:
            raise SystemExit("listings.source is gsheet but no listings.url (or LISTINGS_URL) is set.")
        resp = requests.get(_gsheet_csv_url(url), timeout=30)
        resp.raise_for_status()
        if "<html" in resp.text[:200].lower():
            raise SystemExit(
                "Google returned an HTML page, not CSV — the sheet isn't readable without login.\n"
                "Fix: File → Share → Anyone with the link (Viewer), or File → Share → "
                "Publish to web → CSV, and use that URL."
            )
        raw = _rows_from_csv_text(resp.text)
    elif source == "xlsx":
        raw = _rows_from_xlsx(_resolve_path(lc["path"]))
    elif source == "csv":
        raw = _rows_from_csv_text(_resolve_path(lc["path"]).read_text(encoding="utf-8-sig"))
    else:
        raise SystemExit(f"Unknown listings.source: {source!r} (use csv, xlsx or gsheet)")

    colmap = {field: _norm(header) for field, header in lc["columns"].items() if header}
    listings = []
    for row in raw:
        lookup = {_norm(k): (v or "").strip() for k, v in row.items() if k}
        item = {field: lookup.get(header, "") for field, header in colmap.items()}
        if not item.get("ref"):
            continue  # a row with no ref can't be tracked or linked
        item["_raw"] = row
        listings.append(item)
    return listings


def find_listing(cfg: dict, ref: str) -> dict:
    for item in load_listings(cfg):
        if item["ref"].strip().lower() == ref.strip().lower():
            return item
    raise SystemExit(f"No listing with ref {ref!r} in the sheet.")


def active_listings(cfg: dict) -> list[dict]:
    allowed = {s.strip().lower() for s in cfg["listings"]["active_statuses"]}
    out = []
    for item in load_listings(cfg):
        status = item.get("status", "").strip().lower()
        if not status or status in allowed:
            out.append(item)
    return out


# --------------------------------------------------------------------------- whatsapp

def whatsapp_link(cfg: dict, listing: dict) -> str:
    number = re.sub(r"\D", "", str(cfg["whatsapp"]["number"]))
    if not number or set(number) == {"0"}:
        raise SystemExit("Set a real whatsapp.number in config/agent.yaml (or WHATSAPP_NUMBER in .env).")
    prefill = cfg["whatsapp"]["prefill"].format(
        ref=listing.get("ref", ""), title=listing.get("title", "your listing")
    )
    return f"https://wa.me/{number}?text={urllib.parse.quote(prefill)}"


def cta_line(cfg: dict, listing: dict) -> str:
    return cfg["whatsapp"]["cta_template"].format(link=whatsapp_link(cfg, listing))


# --------------------------------------------------------------------------- state

def _state_path(name: str) -> Path:
    STATE_DIR.mkdir(exist_ok=True)
    return STATE_DIR / f"{name}.json"


def read_state(name: str) -> dict:
    path = _state_path(name)
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return {}


def write_state(name: str, data: dict) -> None:
    _state_path(name).write_text(json.dumps(data, indent=2, sort_keys=True), encoding="utf-8")


def now_utc() -> datetime:
    return datetime.now(timezone.utc)


def days_ago(n: int) -> datetime:
    return now_utc() - timedelta(days=n)


def pick_listings(cfg: dict, count: int) -> list[dict]:
    """Choose the least-recently-posted active listings, respecting the cooldown."""
    posted = read_state("posted").get("listings", {})
    cooldown = int(cfg["schedule"]["cooldown_days"])
    cutoff = days_ago(cooldown)

    candidates = []
    for item in active_listings(cfg):
        last = posted.get(item["ref"], {}).get("last_posted")
        last_dt = datetime.fromisoformat(last) if last else None
        if last_dt and last_dt > cutoff:
            continue  # still cooling down
        candidates.append((last_dt or datetime.min.replace(tzinfo=timezone.utc), item))

    if not candidates:
        return []
    candidates.sort(key=lambda pair: pair[0])
    return [item for _, item in candidates[:count]]


def mark_posted(listing_ref: str, platform: str, post_id: str) -> None:
    state = read_state("posted")
    listings = state.setdefault("listings", {})
    entry = listings.setdefault(listing_ref, {})
    entry["last_posted"] = now_utc().isoformat()
    entry.setdefault("posts", []).append(
        {"platform": platform, "post_id": post_id, "at": now_utc().isoformat()}
    )
    write_state("posted", state)


def recent_posts(cfg: dict) -> list[dict]:
    """Every post the agent made inside the reply lookback window, newest first."""
    cutoff = days_ago(int(cfg["replies"]["lookback_days"]))
    out = []
    for ref, entry in read_state("posted").get("listings", {}).items():
        for post in entry.get("posts", []):
            if datetime.fromisoformat(post["at"]) >= cutoff:
                out.append({**post, "ref": ref})
    out.sort(key=lambda p: p["at"], reverse=True)
    return out
